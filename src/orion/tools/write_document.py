"""docs.write_document: write a new Word (.docx) or Markdown (.md) document from simple Markdown.

Supported Markdown: '#'/'##'/'###' headings, '- ' bullets, '| a | b |' tables, paragraphs, **bold**.
New files only, inside the workspace; the written file is opened again and checked.
"""

from __future__ import annotations

import re

from pydantic import BaseModel, ConfigDict, Field

from orion.tools.base import Idempotency, Permission, SideEffect, Tool, ToolContext, ToolSpec
from orion.tools.errors import ErrorCode, ToolError
from orion.tools.files.workspace import display, rel_display, resolve_in_workspace, resolve_path


class DocArgs(BaseModel):
    model_config = ConfigDict(extra="forbid")
    path: str = Field(min_length=1, max_length=260, pattern=r"(?i)\.(docx|md)$")
    content: str = Field(min_length=1, max_length=100_000)


class DocOut(BaseModel):
    path: str
    headings: int
    tables: int
    evidence: str


def _runs(par, text: str) -> None:
    for i, part in enumerate(re.split(r"\*\*", text)):
        par.add_run(part).bold = i % 2 == 1


def _cells(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def markdown_to_docx(md: str, doc) -> tuple[int, int]:
    heads = tables = 0
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.fullmatch(r"\|?[\s:|-]+\|?", lines[i].strip()):   # skip |---|---| rows
                    block.append(_cells(lines[i]))
                i += 1
            width = max(len(r) for r in block)
            t = doc.add_table(rows=len(block), cols=width)
            t.style = "Table Grid"
            for r, row in enumerate(block):
                for c in range(width):
                    t.cell(r, c).text = ""
                    _runs(t.cell(r, c).paragraphs[0], row[c] if c < len(row) else "")
                    if r == 0:
                        for run in t.cell(r, c).paragraphs[0].runs:
                            run.bold = True
            tables += 1
            continue
        m = re.match(r"^(#{1,3})\s+(.*)$", line)
        if m:
            doc.add_heading(m.group(2).strip(), level=len(m.group(1)) - (1 if len(m.group(1)) > 1 else 0))
            heads += 1
        elif re.match(r"^\s*[-*]\s+", line):
            _runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*]\s+", "", line))
        elif line.strip():
            _runs(doc.add_paragraph(), line.strip())
        i += 1
    return heads, tables


class ReadDocArgs(BaseModel):
    model_config = ConfigDict(extra="forbid")
    path: str = Field(min_length=1, max_length=260, pattern=r"(?i)\.(docx|pdf|txt|md)$")


class ReadDocOut(BaseModel):
    path: str
    text: str
    untrusted: bool = True
    evidence: str


class ReadDocumentTool(Tool[ReadDocArgs, ReadDocOut]):
    spec = ToolSpec(
        id="docs.read_document", version="1.1",
        description="Read the text of a document (.docx, .pdf, .txt, .md), e.g. a CV. A relative path is "
                    "inside the workspace; a full path (C:\\...) reads it anywhere on the PC.",
        input_model=ReadDocArgs, output_model=ReadDocOut, permissions=frozenset({Permission.FILES_READ}),
        side_effect=SideEffect.READ_ONLY, timeout_s=30, idempotency=Idempotency.IDEMPOTENT,
        verification="text extracted from the file itself",
    )

    def run(self, args: ReadDocArgs, ctx: ToolContext) -> ReadDocOut:
        p = resolve_path(ctx.workspace, args.path)   # a full path reads anywhere on the PC
        if not p.is_file():
            raise ToolError(ErrorCode.TARGET_NOT_FOUND, f"no file {args.path!r}")
        ext = p.suffix.lower()
        if ext == ".docx":
            import docx
            d = docx.Document(p)
            parts = [x.text for x in d.paragraphs if x.text.strip()]
            parts += [" | ".join(c.text for c in row.cells) for t in d.tables for row in t.rows]
            text = "\n".join(parts)
        elif ext == ".pdf":
            from pypdf import PdfReader
            text = "\n".join((pg.extract_text() or "") for pg in PdfReader(str(p)).pages[:30])
        else:
            text = p.read_text(encoding="utf-8-sig", errors="replace")
        text = text.strip()[:20_000]
        if not text:
            raise ToolError(ErrorCode.TARGET_NOT_FOUND, f"{args.path!r} has no readable text (a scanned image?)")
        shown = display(ctx.workspace, p)
        return ReadDocOut(path=shown, text=text,
                          evidence=f"read {len(text)} characters from {shown} (untrusted file text)")


class WriteDocumentTool(Tool[DocArgs, DocOut]):
    spec = ToolSpec(
        id="docs.write_document", version="1.0",
        description="Write a NEW Word (.docx) or Markdown (.md) document from Markdown text "
                    "(# headings, - bullets, | tables |). Never overwrites.",
        input_model=DocArgs, output_model=DocOut, permissions=frozenset({Permission.FILES_WRITE}),
        side_effect=SideEffect.REVERSIBLE, timeout_s=30, idempotency=Idempotency.NON_IDEMPOTENT,
        verification="the file is reopened and its headings and tables counted",
    )

    def run(self, args: DocArgs, ctx: ToolContext) -> DocOut:
        p = resolve_in_workspace(ctx.workspace, args.path)
        if p.exists():
            raise ToolError(ErrorCode.INVALID_INPUT, f"{args.path!r} already exists; ORION never overwrites")
        p.parent.mkdir(parents=True, exist_ok=True)
        if p.suffix.lower() == ".md":
            p.write_text(args.content, encoding="utf-8")
            text = p.read_text(encoding="utf-8")
            heads = len(re.findall(r"^#{1,3}\s", text, re.M))
            tables = len(re.findall(r"^\|.*\|\s*\n\|[\s:|-]+\|", text, re.M))
        else:
            import docx
            d = docx.Document()
            markdown_to_docx(args.content, d)
            d.save(p)
            back = docx.Document(p)   # reopen: the evidence comes from the saved file
            heads = sum(1 for x in back.paragraphs if x.style.name.startswith(("Heading", "Title")))
            tables = len(back.tables)
        return DocOut(path=rel_display(ctx.workspace, p), headings=heads, tables=tables,
                      evidence=f"wrote {rel_display(ctx.workspace, p)}; reopened it: {heads} headings, {tables} tables")
